"""Rebuild synthetic binaries using `uv sync --group corpus` first.

Source facts and ACLs live in demo_corpus/source.json; rebuilding updates the manifest hashes.
This is authoring tooling, not the production ingestion/parser implementation.
"""

import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "demo_corpus"
NOTICE = "Contenido ficticio de Lumen Demo para pruebas de DocuMind AI. No es una política real."


def author_document(item: dict[str, Any]) -> None:
    target = CORPUS / "files" / item["file"]
    sections = item["sections"]
    if target.suffix == ".docx":
        doc = Document()
        page = doc.sections[0]
        page.top_margin = page.bottom_margin = Inches(0.8)
        normal = doc.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
        for border in doc.styles["Title"].element.iter(qn("w:pBdr")):
            border.getparent().remove(border)
        doc.add_paragraph(item["title"], "Title")
        doc.add_paragraph(NOTICE)
        for section in sections:
            doc.add_heading(section["heading"], level=1)
            doc.add_paragraph(section["text"])
        doc.core_properties.author = "DocuMind Demo"
        doc.core_properties.created = doc.core_properties.modified = datetime(
            2026, 1, 1, tzinfo=UTC
        )
        doc.save(str(target))
    elif target.suffix == ".pdf":
        styles = getSampleStyleSheet()
        flow: list[Any] = [Paragraph(escape(item["title"]), styles["Title"])]
        flow.extend([Paragraph(NOTICE, styles["BodyText"]), Spacer(1, 18)])
        for section in sections:
            flow.extend(
                [
                    Paragraph(escape(section["heading"]), styles["Heading2"]),
                    Paragraph(escape(section["text"]), styles["BodyText"]),
                    Spacer(1, 12),
                ]
            )
        SimpleDocTemplate(
            str(target), title=item["title"], author="DocuMind Demo", invariant=1
        ).build(flow)
    else:
        chunks = [f"# {item['title']}", NOTICE]
        for section in sections:
            chunks.extend([f"## {section['heading']}", section["text"]])
        target.write_text("\n\n".join(chunks) + "\n", encoding="utf-8", newline="\n")


def main() -> None:
    source = json.loads((CORPUS / "source.json").read_text(encoding="utf-8"))
    (CORPUS / "files").mkdir(exist_ok=True)
    documents = []
    for item in source["documents"]:
        author_document(item)
        content = (CORPUS / "files" / item["file"]).read_bytes()
        entry = {key: value for key, value in item.items() if key != "sections"}
        entry.update(sha256=hashlib.sha256(content).hexdigest(), size_bytes=len(content), version=1)
        documents.append(entry)
    manifest = {"version": source["version"], "synthetic": True, "documents": documents}
    (CORPUS / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"PASS: {len(documents)} synthetic documents and SHA-256 manifest")


if __name__ == "__main__":
    main()
